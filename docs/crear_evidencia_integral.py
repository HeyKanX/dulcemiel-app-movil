from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\thehe\Documents\SolucionesWeb\Informe\Informe Semana 4 Dulce Miel APA 7.docx")
OUTPUT = Path(r"C:\Users\thehe\Documents\SolucionesWeb\Informe\Evidencia Semana 4 Proyecto Dulce Miel.docx")


def font(run, size=12, bold=None, italic=None, color="000000"):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Times New Roman")
    rpr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(text), bold=True)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5)
    font(p.add_run(text))
    return p


def lead(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5)
    font(p.add_run(label + ". "), bold=True)
    font(p.add_run(text))
    return p


def insert_elements_before(anchor, elements):
    for element in elements:
        anchor._p.addprevious(element)


def capture_new_elements(doc, builder):
    body_el = doc._element.body
    marker = doc.add_paragraph("__TEMP_INSERTION_MARKER__")._p
    builder()
    added = []
    current = marker.getnext()
    while current is not None and current.tag != qn("w:sectPr"):
        added.append(current)
        current = current.getnext()
    copies = [deepcopy(element) for element in added]
    for element in added:
        body_el.remove(element)
    body_el.remove(marker)
    return copies


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def margins(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def borders(table):
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "D9D9D9")
        tbl_borders.append(node)
    table._tbl.tblPr.append(tbl_borders)


def add_stack_table(doc):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    font(p.add_run("Tabla 5"), bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run("Stack técnico y responsabilidad prevista"), italic=True)
    rows = [
        ("Android Studio", "Entorno para crear, ejecutar, depurar y construir la aplicación Android."),
        ("Kotlin", "Lenguaje principal para la lógica, modelos y componentes de la aplicación."),
        ("Jetpack Compose", "Construcción declarativa de las pantallas del prototipo y sus estados visuales."),
        ("MVVM", "Separación entre interfaz, estado de presentación mediante ViewModel y acceso a datos."),
        ("API REST", "Integración futura con autenticación, catálogo, carrito y pedidos del sistema existente."),
        ("Firebase", "Notificaciones autorizadas con Cloud Messaging y registro técnico de fallos con Crashlytics."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for i, value in enumerate(("Tecnología", "Responsabilidad")):
        cell = table.rows[0].cells[i]
        cell.width = Inches((1.6, 4.9)[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell)
        shade(cell, "365F91")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(value), 10, bold=True, color="FFFFFF")
    hdr = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    hdr.append(marker)
    for ri, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches((1.6, 4.9)[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if ri % 2:
                shade(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            font(p.add_run(value), 10)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    font(p.add_run("Nota. "), 10, italic=True)
    font(p.add_run("Elaboración propia a partir de la arquitectura definida para la migración móvil."), 10)


doc = Document(SOURCE)

# Actualiza portada y propiedades.
title = next(p for p in doc.paragraphs if p.text.startswith("Planificación y prototipo inicial"))
title.text = "Artefactos y evidencia técnica del proyecto móvil Dulce Miel"
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.style = doc.styles["Title"]
font(title.runs[0], bold=True)
for p in doc.paragraphs:
    if p.text == "Semana 4":
        p.text = "Evidencia de la semana 4"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.runs[0])
        break

objective_anchor = next(p for p in doc.paragraphs if p.text == "Objetivo del avance")

def build_context():
    heading(doc, "Contexto de la empresa")
    body(doc, "Dulce Miel es una organización apícola ubicada en La Libertad, Perú, dedicada a la producción y comercialización de miel y productos derivados para clientes minoristas y mayoristas. El proyecto web previo centralizó la consulta del catálogo, la autenticación, el carrito, las compras, el seguimiento de pedidos y la administración comercial. Esta solución emplea un frontend web, una API REST y una base de datos relacional desplegada en Azure, por lo que constituye la base funcional y técnica de la migración móvil.")
    body(doc, "La aplicación Android se orienta inicialmente al cliente. Los empleados y administradores continuarán utilizando el panel web para evitar duplicar funciones y mantener un alcance académico realista. El equipo está integrado por Erick Diego Enriquez Pacheco y Cosme Vela Horna, quienes participan en análisis, diseño, implementación, pruebas y documentación.")
    heading(doc, "Análisis del problema")
    body(doc, "El problema seleccionado es la dependencia de una experiencia basada en navegador para consultar productos, realizar pedidos y revisar su estado. Aunque la plataforma web puede adaptarse a pantallas pequeñas, no ofrece una interacción nativa centrada en Android ni aprovecha directamente la persistencia de sesión, las notificaciones del dispositivo y una navegación táctil diseñada para consultas frecuentes.")
    body(doc, "También existe una oportunidad de mejorar la calidad de las decisiones de compra. La aplicación debe mostrar precios y existencias antes de añadir productos, evitar selecciones automáticas y brindar información verificable sobre conservación y uso. Estas decisiones reducen compras accidentales, fortalecen la confianza y vinculan el proyecto con el ODS 12.")
    heading(doc, "Antecedentes")
    body(doc, "El antecedente principal es la plataforma web Dulce Miel desarrollada para atender la ausencia de un canal digital integrado. Su API REST y Azure SQL conservarán la lógica y los datos comerciales. La propuesta de la semana 2 definió usuarios, alcance y stack; la semana 3 incorporó restricciones de conectividad, seguridad, privacidad, accesibilidad, rendimiento y uso responsable de recursos, además de criterios éticos para las campañas digitales.")
    body(doc, "La migración será incremental y no una reconstrucción completa. La aplicación consumirá servicios existentes y Firebase cumplirá funciones complementarias. Esta decisión reduce duplicidades, permite reutilizar reglas de negocio ya establecidas y concentra el desarrollo inicial en la experiencia móvil del cliente.")
    heading(doc, "Restricciones realistas")
    lead(doc, "Conectividad", "La red puede ser lenta o intermitente. El catálogo y el historial reciente deberán consultarse desde una caché identificada con fecha de actualización. Un pedido solo se considerará registrado después de la confirmación del servidor y los reintentos no deben duplicarlo.")
    lead(doc, "Seguridad", "Toda comunicación se realizará mediante HTTPS. Los tokens tendrán vencimiento y almacenamiento seguro; la autorización se validará en el servidor. Contraseñas, tokens y datos personales no se incluirán en registros técnicos.")
    lead(doc, "Privacidad", "La aplicación solicitará únicamente los datos y permisos necesarios. No requiere contactos, micrófono, cámara ni ubicación permanente para su alcance inicial. Las notificaciones promocionales dependerán del consentimiento configurable.")
    lead(doc, "Accesibilidad", "Los controles principales tendrán etiquetas para lector de pantalla, áreas táctiles adecuadas y texto escalable. Los estados no dependerán exclusivamente del color y los flujos críticos se probarán con TalkBack.")
    lead(doc, "Rendimiento", "Las listas usarán carga gradual, las imágenes se optimizarán y las tareas de red no bloquearán la interfaz. Se observarán tiempos de inicio, cierres inesperados, consumo de red y tamaño de caché.")
    lead(doc, "Uso responsable de recursos", "La aplicación evitará sincronizaciones continuas, datos temporales ilimitados y notificaciones repetitivas. Tampoco buscará aumentar el tiempo de pantalla ni inducir compras innecesarias.")


insert_elements_before(objective_anchor, capture_new_elements(doc, build_context))

prototype_anchor = next(p for p in doc.paragraphs if p.text == "Prototipo inicial")

def build_annex_and_evidence():
    heading(doc, "Anexo 1 Casos de marketing irresponsable en aplicaciones")
    body(doc, "Las campañas dentro de una aplicación incluyen anuncios, banners, promociones, recomendaciones, notificaciones y decisiones predeterminadas. Se consideran irresponsables cuando engañan, presionan o dificultan una elección libre. Los casos siguientes resumen prácticas documentadas por la OCDE, la Comisión Federal de Comercio de Estados Unidos e Indecopi, y proponen su tratamiento ético en Dulce Miel.")
    lead(doc, "Caso 1 Urgencia y escasez falsas", "Algunas campañas muestran temporizadores que se reinician o mensajes de pocas unidades sin respaldo. La mala práctica aprovecha la presión temporal para reducir la reflexión. Dulce Miel solo mostrará vigencia y stock obtenidos de datos confirmados, sin contadores simulados.")
    lead(doc, "Caso 2 Productos o costos preseleccionados", "La incorporación automática de artículos, servicios o cargos puede producir compras no deseadas. En la propuesta ética, cada producto requiere una acción afirmativa, el carrito comienza sin extras y el total cambia únicamente después de una acción visible.")
    lead(doc, "Caso 3 Notificaciones promocionales excesivas", "El envío frecuente, en horarios inadecuados o con mensajes de culpa convierte la notificación en presión comercial. La alternativa separa los avisos transaccionales de las promociones, solicita consentimiento por categorías y permite desactivar publicidad sin perder información del pedido.")
    lead(doc, "Caso 4 Recolección forzada e interferencia visual", "Solicitar datos o permisos innecesarios y resaltar aceptar mientras se oculta rechazar limita la autonomía. Dulce Miel utilizará opciones comprensibles y visibles, solicitará solo datos relacionados con la compra y mantendrá la función principal cuando se rechacen permisos opcionales.")
    heading(doc, "Propuestas éticas")
    for statement in (
        "Toda afirmación sobre precio, disponibilidad, origen, calidad o sostenibilidad debe poder comprobarse.",
        "El precio total y las condiciones relevantes se muestran antes de confirmar el pedido.",
        "Ningún producto, pago adicional, permiso o suscripción se activa de forma predeterminada.",
        "Aceptar, rechazar, volver y cancelar deben presentarse con lenguaje claro y visibilidad equivalente.",
        "Las notificaciones promocionales requieren autorización y se gestionan por separado de los avisos del pedido.",
        "Las campañas vinculadas al ODS 12 deben facilitar decisiones responsables y evitar afirmaciones ambientales vagas.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = 2
        font(p.add_run(statement))
    heading(doc, "Evidencia técnica")
    body(doc, "El prototipo inicial fue desarrollado en Android Studio con Kotlin y Jetpack Compose. Mantiene el estado de navegación, la selección de productos y el carrito mediante ViewModel, como primera aproximación a MVVM. Se compiló correctamente, se instaló y se ejecutó en un dispositivo virtual Pixel 10 con Android 16. Los datos son simulados y permiten validar el recorrido antes de integrar la API REST y Firebase.")
    add_stack_table(doc)


insert_elements_before(prototype_anchor, capture_new_elements(doc, build_annex_and_evidence))

# Amplía la introducción para reflejar la evidencia integral.
intro = next(p for p in doc.paragraphs if p.text.startswith("El presente documento organiza"))
intro.text = ("El presente documento reúne la evidencia de la semana 4 del proyecto móvil Dulce Miel. "
              "Desarrolla el contexto de la empresa, el análisis del problema, los antecedentes, las restricciones realistas, "
              "el alcance, las reglas de negocio y el Product Backlog. Además, incorpora un anexo sobre campañas de marketing "
              "irresponsable con propuestas éticas, y presenta la evidencia técnica del prototipo Android desarrollado con Kotlin.")
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
intro.paragraph_format.line_spacing = 2
intro.paragraph_format.first_line_indent = Inches(0.5)
for run in intro.runs:
    font(run)

# Referencias adicionales del anexo.
for ref in (
    "Federal Trade Commission. (2022). Bringing dark patterns to light. https://www.ftc.gov/reports/bringing-dark-patterns-light",
    "Instituto Nacional de Defensa de la Competencia y de la Protección de la Propiedad Intelectual. (2022, 28 de marzo). Cyber Days: Indecopi advierte que proveedores podrían usar patrones oscuros en sus páginas web para influir en decisiones de compra. https://www.gob.pe/institucion/indecopi/noticias/595426",
    "Organización para la Cooperación y el Desarrollo Económicos. (2022). Dark commercial patterns (OECD Digital Economy Papers No. 336). OECD Publishing. https://doi.org/10.1787/44f5e846-en",
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 2
    font(p.add_run(ref))

doc.core_properties.title = "Artefactos y evidencia técnica del proyecto móvil Dulce Miel"
doc.core_properties.subject = "Evidencia integral de la semana 4"
doc.core_properties.author = "Erick Diego Enriquez Pacheco y Cosme Vela Horna"
doc.save(OUTPUT)
print(OUTPUT)
